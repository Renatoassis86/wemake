import sys, json, os, re, warnings
warnings.filterwarnings('ignore')
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
SRC = r'C:\repositorio\wemake\projetos_wemake\app_comercial_We Make\Academia_wemake'
OUT = r'C:\repositorio\wemake\projetos_wemake\app_comercial_We Make\src\content\academia'

def P(x): return {'t': 'p', 'x': x}
def H(t, x): return {'t': t, 'x': x}
def cell(x): return {'b': [{'t': 'p', 'x': x}] if x not in (None, '') else []}
def table(rows, kind=None):
    t = {'t': 'table', 'rows': [[cell(c) for c in r] for r in rows]}
    if kind: t['kind'] = kind
    return t
def s(v):
    if v is None: return ''
    if isinstance(v, float) and v == int(v): v = int(v)
    return str(v).strip()
def money(v):
    try: return 'R$ ' + f'{float(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    except Exception: return s(v)

# ---------- Diagnóstico ----------
wb = openpyxl.load_workbook(os.path.join(SRC, 'Docs Diagramados', '3.1 Diagnostico_Espaco_Maker_We_Make_2027.xlsx'))
B = []
ws = wb['Orientações']
B += [H('title', s(ws['A1'].value)), P(s(ws['A2'].value)), H('h1', 'Como utilizar esta planilha')]
for r in range(5, 11): B.append({'t': 'li', 'x': s(ws.cell(r, 2).value)})
B.append(H('h2', 'Fluxo do diagnóstico'))
B.append(table([[s(ws.cell(r, 1).value), s(ws.cell(r, 2).value)] for r in range(13, 18)]))
B.append(P('**' + s(ws['A19'].value) + '**'))
B.append(H('h2', 'Legenda de cores'))
B.append({'t': 'li', 'x': s(ws['B22'].value)})
for r in range(23, 27):
    if ws.cell(r, 2).value: B.append({'t': 'li', 'x': s(ws.cell(r, 2).value)})
ws = wb['Sala e Evidências']
B.append(H('h1', 'Sala maker e evidências'))
B.append(P(s(ws['A2'].value)))
B.append(H('h2', '1. Caracterização do ambiente — preenchimento da escola'))
B.append(table([['Pergunta', 'Resposta da escola']] + [[s(ws.cell(r, 1).value), ''] for r in range(5, 10)]))
B.append(H('h2', '2. Checklist de medidas e briefing — preenchimento da escola'))
B.append(table([[s(ws.cell(12, c).value) for c in range(1, 5)]] + [[s(ws.cell(r, 1).value), s(ws.cell(r, 2).value), '', ''] for r in range(13, 33)]))
B.append(H('h2', '3. Evidências obrigatórias — preenchimento da escola'))
B.append(table([[s(ws.cell(35, c).value) for c in range(1, 5)]] + [[s(ws.cell(r, 1).value), '', '', ''] for r in range(36, 47)]))
B.append(H('h2', '4. Parecer técnico do ambiente — preenchimento da We Make'))
B.append(table([[s(ws.cell(50, c).value) for c in range(1, 4)]] + [[s(ws.cell(r, 1).value), '', ''] for r in range(51, 62)]))
B.append(P('**' + s(ws['A63'].value) + '**'))
for name in ('Recursos Reutilizáveis', 'Recursos Consumíveis'):
    ws = wb[name]
    B.append(H('h1', name))
    B.append(P('Ciano-claro = preenchimento da escola · Marfim = preenchimento da We Make · Quantidades de aquisição são calculadas automaticamente. A planilha traz, para cada item, a especificação e a quantidade recomendada. As colunas de preenchimento (Escola possui?, Qtd. existente, Marca / modelo / observação, Parecer We Make, Qtd. aproveitável, Qtd. a adquirir, Ação recomendada e Custo estimado da aquisição) ficam em branco até o diagnóstico.'))
    rows = [['Categoria', 'Item', 'Especificação / descrição', 'Qtd. recomendada', 'Unid.', 'Valor ref. unit.', 'Referência']]
    for r in range(5, ws.max_row + 1):
        if not ws.cell(r, 2).value: continue
        v = ws.cell(r, 7).value
        rows.append([s(ws.cell(r, 1).value), s(ws.cell(r, 2).value), s(ws.cell(r, 3).value), s(ws.cell(r, 4).value), s(ws.cell(r, 5).value), money(v) if v not in (None, '') else '', s(ws.cell(r, 6).value)])
    B.append(table(rows, 'resources'))
ws = wb['Resumo We Make']
B.append(H('h1', 'Resumo We Make'))
B.append(P(s(ws['A2'].value)))
B.append(table([['Indicador', 'Valor']] + [[s(ws.cell(r, 1).value), 'calculado automaticamente'] for r in range(5, 12)]))
B.append(H('h2', 'Parecer geral do Espaço Maker'))
B.append(table([['Item', 'Parecer']] + [[s(ws.cell(r, 1).value), ''] for r in range(15, 19)]))
B.append(P('**' + s(ws['A20'].value).capitalize() + '**'))
ws = wb['Listas']
B.append(H('h1', 'Listas de valores'))
B.append(table([[s(ws.cell(1, c).value) for c in (1, 2, 3)]] + [[s(ws.cell(r, c).value) for c in (1, 2, 3)] for r in range(2, 8)]))
json.dump({'slug': 'planilha-diagnostico', 'source': '3.1 Diagnostico_Espaco_Maker_We_Make_2027.xlsx', 'blocks': B}, open(os.path.join(OUT, 'planilha-diagnostico.json'), 'w', encoding='utf-8'), ensure_ascii=False)

# ---------- Painel ----------
wb = openpyxl.load_workbook(os.path.join(SRC, 'Docs Diagramados', '8. Painel_Mestre_Implantacao_We_Make_2027.xlsx'))
B = []
ws = wb['Como usar']
B += [H('title', s(ws['A1'].value)), P(s(ws['A3'].value))]
r = 6
maxr = ws.max_row
while r <= maxr:
    a = s(ws.cell(r, 1).value)
    if re.match(r'^\d\. ', a) and len(a) < 80 and not ws.cell(r, 2).value:
        B.append(H('h1', a[3:].capitalize() if a.isupper() else a[3:]))
        r += 1
        while r <= maxr and ws.cell(r, 1).value and not ws.cell(r, 2).value and not (re.match(r'^\d\. ', s(ws.cell(r, 1).value)) and len(s(ws.cell(r, 1).value)) < 80):
            txt = s(ws.cell(r, 1).value)
            if re.match(r'^\d\. ', txt) and re.search(r' / \d\. ', txt):
                for part in re.split(r' / (?=\d\. )', txt): B.append({'t': 'li', 'x': re.sub(r'^\d\. ', '', part.strip())})
            else:
                B.append(P(txt))
            r += 1
        if r <= maxr and ws.cell(r, 2).value:
            rows = []
            while r <= maxr and ws.cell(r, 2).value:
                rows.append([s(ws.cell(r, c).value) for c in range(1, 6)])
                r += 1
            w = max(max((i + 1 for i, c in enumerate(row) if c != ''), default=0) for row in rows)
            rows = [row[:w] for row in rows]
            B.append(table(rows))
        continue
    r += 1
ws = wb['Matriz documentos']
B.append(H('h1', 'Matriz de documentos da implantação'))
B.append(P(s(ws['A3'].value)))
B.append(table([[s(ws.cell(6, c).value) for c in range(1, 9)]] + [[s(ws.cell(r, c).value) for c in range(1, 9)] for r in range(7, 16)]))
B.append(H('h2', 'Macroetapas × marcos do Painel'))
t17 = s(ws['A17'].value)
B.append(P(t17.split(' / ', 1)[1] if ' / ' in t17 else t17))
B.append(H('h2', 'Padrão de arquivamento recomendado'))
for r in range(22, 26): B.append({'t': 'li', 'x': s(ws.cell(r, 2).value)})
ws = wb['Painel']
B.append(H('h1', 'Estrutura das abas do Painel'))
B.append(H('h2', 'Aba Painel'))
B.append(P(s(ws['A1'].value).capitalize() + '. ' + s(ws['A3'].value)))
B.append(P('Indicadores de topo: Escolas no painel · Go-Live concluído · Com risco alto/crítico · Prazos vencidos.'))
B.append(table([['Etapa', 'Concluídas', 'Total', '% concluído']] + [[s(ws.cell(r, 1).value), 'automático', 'automático', 'automático'] for r in range(11, 20)]))
B.append(table([['Risco', 'Escolas', '%']] + [[s(ws.cell(r, 6).value), 'automático', 'automático'] for r in range(11, 15)]))
B.append(H('h3', 'Escolas que exigem atenção'))
B.append(P('Colunas: ' + ', '.join(s(ws.cell(23, c).value) for c in range(1, 9)) + '.'))
B.append(P(s(ws['A24'].value)))
ws = wb['Implantacao']
B.append(H('h2', 'Aba Implantacao'))
B.append(P('Base operacional principal: uma linha por escola. Colunas, na ordem:'))
for c in range(1, 21):
    v = s(ws.cell(1, c).value)
    if v: B.append({'t': 'li', 'x': v})
B.append(H('h3', 'Como usar (na própria aba)'))
for r in range(2, 7): B.append({'t': 'li', 'x': s(ws.cell(r, 23).value)})
ws = wb['Listas']
B.append(H('h2', 'Aba Listas'))
B.append(table([[s(ws.cell(1, c).value) for c in (1, 2, 3, 5)]] + [[s(ws.cell(r, c).value) for c in (1, 2, 3, 5)] for r in range(2, 11)]))
json.dump({'slug': 'painel-mestre', 'source': '8. Painel_Mestre_Implantacao_We_Make_2027.xlsx', 'blocks': B}, open(os.path.join(OUT, 'painel-mestre.json'), 'w', encoding='utf-8'), ensure_ascii=False)
print('ok')
